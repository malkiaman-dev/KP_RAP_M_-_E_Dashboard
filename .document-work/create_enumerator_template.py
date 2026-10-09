from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path('Data Quality Sessions/Enumerator_Error_aur_Hal_Template_Roman_Urdu_07-Oct-2026.docx')
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.65)
normal = d.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name in ['Title', 'Heading 1', 'Heading 2']:
    d.styles[name].font.name = 'Calibri'
    d.styles[name].font.color.rgb = RGBColor(0, 0, 0)
d.styles['Title'].font.size = Pt(21)
d.styles['Title'].paragraph_format.space_after = Pt(8)
d.styles['Heading 1'].font.size = Pt(12)
d.styles['Heading 1'].paragraph_format.space_before = Pt(12)
d.styles['Heading 1'].paragraph_format.space_after = Pt(6)
p = d.add_paragraph('KP-RAP PROJECT')
p.runs[0].bold = True
p.runs[0].font.size = Pt(10)
d.add_paragraph('Enumerator Errors aur Hal ka Form', 'Title')
d.add_paragraph('D.I. Khan aur Hangu | Data Quality Session | 07 October 2026')
d.add_paragraph('Session mein discuss hone wale apne errors is form mein likhein. Har error ke saamne us ka hal aur steps likhein, taake error durust ho aur dobara na ho.')

def field(label, value):
    p = d.add_paragraph()
    p.add_run(label + '  ').bold = True
    p.add_run(value)
    p.paragraph_format.space_after = Pt(9)

d.add_paragraph('Enumerator ki details', 'Heading 1')
field('Enumerator ka naam', '_' * 53)
field('Enumerator ID ya username', '_' * 43)
field('District', 'D.I. Khan / Hangu     Village ya ilaqa  ____________________')
field('Form bharne ki date', '____________     Supervisor ka naam  ________________')
d.add_paragraph('Mere errors aur un ka hal', 'Heading 1')
d.add_paragraph('Har error alag row mein likhein. Jahan zaroori ho, survey type aur record ID bhi dein. Hal wale column mein batayein ke aap kya steps lenge. Zaroorat ho to mazeed rows add karein.')
t = d.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [.42, 2.65, 3.90]
headers = ['No.', 'Error ya ghalti kya thi', 'Hal aur main is par kaise amal karunga / karungi']
for i, (c, text) in enumerate(zip(t.rows[0].cells, headers)):
    c.width = Inches(widths[i])
    c.text = text
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EDF3'); c._tc.get_or_add_tcPr().append(shade)
    for r in c.paragraphs[0].runs: r.bold = True
repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
for n in range(1, 7):
    row = t.add_row()
    row.height = Inches(.53)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for i, c in enumerate(row.cells): c.width = Inches(widths[i])
    row.cells[0].text = str(n)
for row in t.rows:
    no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    for i, c in enumerate(row.cells):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        props = c._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e = OxmlElement('w:' + edge)
            for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:' + key), val)
            borders.append(e)
        props.append(borders)
        margins = OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            e = OxmlElement('w:' + edge); e.set(qn('w:w'), '100'); e.set(qn('w:type'), 'dxa'); margins.append(e)
        props.append(margins)
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.size = Pt(10.5)
p = d.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.add_run('Misal: ').bold = True
p.add_run('Error: Consent screen jaldi skip ki. Hal: Poora consent parh kar sunaunga / sunaungi, sawal ka mauqa dunga / dungi aur respondent ka faisla record kar ke aage barhunga / barhungi.')
p.runs[-1].font.size = Pt(10)
d.add_paragraph('Supervisor ka review', 'Heading 1')
field('Comments ya kis madad ki zaroorat hai', '_' * 32)
field('Enumerator ke dastakhat', '____________   Supervisor ke dastakhat  ____________')
d.core_properties.title = 'Enumerator Errors aur Hal ka Form'
d.core_properties.subject = 'D.I. Khan and Hangu data quality session on 07 October 2026'
d.core_properties.author = 'KP-RAP'
for element in d.element.xpath('//w:pBdr'):
    element.getparent().remove(element)
for style in d.styles:
    for element in style.element.xpath('.//w:pBdr'):
        element.getparent().remove(element)
d.save(out)
print(out.resolve())
