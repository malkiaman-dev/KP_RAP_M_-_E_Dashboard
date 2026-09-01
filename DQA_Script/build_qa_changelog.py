"""Simple note: World Bank comments and the changes made in Version 2.0."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_qa_workplan import (
    COVER,
    GOLD,
    INK,
    MUTED,
    ROOT,
    TEAL,
    WHITE,
    add_page_number,
    apply_palette,
    body,
    cell_margins,
    h1,
    meta_row,
    no_borders,
    P,
    set_run,
    set_row_height,
    shade,
    tbl,
)

OUT_NAME = "RAP_Quality_Assurance_Workplan_2.O_19_August2026_Changes.docx"

ROWS = [
    [
        "Umair Kiani: Add a sign-off block on the cover with dated signatures from AoE and PIU. Do not use the protocol for enumerator orientation until that joint sign-off is in place. The protocol was part of the original ToR and proposal.",
        "Cover now has an AoE and PIU sign-off table (name, designation, signature, date). The protocol says orientation will start only after both have signed. It is stated as part of the original ToR and proposal.",
    ],
    [
        "Umair Kiani: Field supervision is missing (accompaniment, spot checks, unannounced visits). This should be done by someone who is not part of the team collecting survey data.",
        "A field-supervision section was added. The Supervisor (Male) is not part of Household or Girls data collection. He will do accompaniment, spot checks and unannounced visits.",
    ],
    [
        "Umair Kiani: Some surveys are still left in Torghar even if it is paused.",
        "Coverage now says D.I. Khan, Hangu, Lakki Marwat, and remaining work in Torghar (currently paused).",
    ],
    [
        "Umair Kiani: The consent checks only test whether values are consistent, not whether consent was actually carried out. In Torghar, consent screens were tapped through. Add that check.",
        "Consent screens tapped through are now flagged using SurveyCTO speed warnings. AoE will treat this as an integrity finding (Track 2) and will resurvey if consent was not actually obtained.",
    ],
    [
        "Hijab Waheed: Add the protocol AoE will follow when duplicates are found: inform the World Bank team after investigating, and say which submission ID to retain.",
        "AoE will investigate every duplicate, will inform the World Bank team with the details, and will recommend which submission ID to retain. The log names the KEY to retain.",
    ],
    [
        "Umair Kiani: In Torghar, completed households were re-opened around midnight. Add a preventive control: once a household is marked complete on the server, lock it unless PIU/IE documents approval.",
        "AoE will keep a completed household or girl locked on the server. It will be re-opened only with documented PIU / IE approval. Re-entry without approval will be invalidated.",
    ],
    [
        "Hijab Waheed: Telephonic interviews are not allowed in any case.",
        "AoE will not allow telephonic interviews. Household and Girls interviews will be completed in person.",
    ],
    [
        "Hijab Waheed: Add extremely small households too.",
        "Extremely small household size is now a Household roster check, along with extremely large households.",
    ],
    [
        "Hijab Waheed: Share the protocol AoE will follow when such roster cases are detected.",
        "The script flags the case. AoE will verify the roster. If the listed girl was left out, AoE will investigate and will resurvey.",
    ],
    [
        "Hijab Waheed: The listed girl should be the first entry in the siblings roster, and the spelling should match girl_label.",
        "Two checks were added: listed girl not first on the siblings roster, and listed-girl spelling does not match girl_label.",
    ],
    [
        "Umair Kiani: A district-boundary GPS test is weak. A fabricated form completed anywhere in the district would pass it.",
        "GPS checks now also flag missing GPS, points far from other interviews in the same village, GPS jumps during the interview, and one location stored under many village names. AoE will resurvey where GPS is inconsistent.",
    ],
    [
        "Hijab Waheed: Map coordinates to administrative layers and spot-check whether points fall in a residential or remote area.",
        "The script cannot use shapefile administrative layers. It uses the district box and the village cluster already in the export. The Supervisor (Male) will spot-check that GPS is consistent with a residential location.",
    ],
    [
        "Umair Kiani: State the actual minimum duration, not “implausibly short”. In Torghar many surveys were finished in under 15 minutes. Also add module-level duration (consent, roster, harassment, and so on).",
        "The 15-minute floor is now stated for Household and Girls. A form under 15 minutes is an integrity finding. Module-level seconds cannot be run from the export (text-audit files are not in the file). AoE will use total duration and SurveyCTO speed warnings instead. This limit is explained in the cannot-check section.",
    ],
    [
        "Hijab Waheed: Check automated start/end time against manually entered start/end times.",
        "This cannot be run from the export. There is no typed start/end pair in the file, only SurveyCTO starttime and endtime. The script still checks those device times. This is explained in the cannot-check section.",
    ],
    [
        "Hijab Waheed: Add checks for education-expenditure outliers and don’t-know / refuse use by enumerators.",
        "Education expenditure outlier and high don’t-know / refuse by enumerator were added to the daily checks.",
    ],
    [
        "Hijab Waheed: Also flag dummy primary phone numbers.",
        "Dummy primary phone number was added (Household and Girls), along with alternative and neighbour numbers already in place.",
    ],
    [
        "Umair Kiani: Require a photograph of the completed paper test, with QA reviewing a sample; add enumerator-level score distribution checks; add a minimum duration for the assessment module.",
        "Enumerators must photograph the completed test paper and upload it. The script flags a missing file and unusual enumerator score distributions. Assessment-module duration cannot be split from the export. The Supervisor (Male) will review a sample of photographs in the field. Total Girls duration under 15 minutes is treated as integrity.",
    ],
    [
        "Hijab Waheed: Do daily spot checks on whether the test-paper image is correct, blank or random.",
        "The script cannot open the image. The Supervisor (Male) will look at a sample of photographs during spot checks and accompaniment to see whether the page is a real completed test.",
    ],
    [
        "Hijab Waheed: Check Household GPS against Girls GPS for the same girl ID, to confirm both surveys were at the same place.",
        "Girls GPS vs Household GPS for the same girl ID is now a daily check. AoE will resurvey if they do not match.",
    ],
    [
        "Umair Kiani: Every enumerator should keep tablet location on. Restate this in the debrief call.",
        "AoE will require enumerators to keep tablet location on. This will be restated in the debrief call. Missing GPS is flagged.",
    ],
    [
        "Umair Kiani: State the Girls duration threshold. A Girls form under 15 minutes should be invalid and sent to the integrity track, not treated as a routine timing flag.",
        "Girls duration under 15 minutes is now an integrity finding. AoE will suspend pending investigation and will resurvey.",
    ],
    [
        "Umair Kiani: Split response into two tracks. Track 1 can stay as notify-call-guide. Track 2 (integrity) needs immediate suspension, full workload review, invalidation of forms that cannot be verified, and written notice to PIU within 24 hours. An enumerator explanation must not close an integrity finding.",
        "Response is now two tracks. Track 1: notify, call, explain, guide. Track 2: suspend, review all of that enumerator’s forms, invalidate what cannot be verified, notify PIU in writing within 24 hours. Only independent verification closes Track 2.",
    ],
    [
        "Hijab Waheed: State how AoE will handle serious findings. Resurvey when the listed girl is omitted, GPS shows a different area, or the survey was rushed.",
        "A separate resurvey subsection was added. AoE will resurvey for listed girl omitted, GPS in a different area, and survey rushed (under 15 minutes).",
    ],
    [
        "Umair Kiani: Also resurvey for learning-test inconsistencies. Tell enumerators that GPS is captured at intervals without internet, and that inconsistency will mean a repeat survey. List all resurvey triggers in a separate subsection.",
        "Learning-test inconsistency is a resurvey trigger. Enumerators will be told that GPS is captured at intervals without internet. All resurvey cases are listed in the resurvey subsection, including consent not obtained and completed-case re-entry.",
    ],
]


def setup_header(doc) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(header.add_run(), "KP-RAP  ", size=8, bold=True, color=TEAL)
    set_run(header.add_run(), "·  World Bank comments and changes", size=8, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(
        footer.add_run(),
        "KP-RAP  ·  Comments and changes  ·  19 August 2026  ·  Page ",
        size=8,
        color=MUTED,
    )
    add_page_number(footer)

    fp = section.first_page_footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(
        fp.add_run(),
        "KP-RAP  ·  World Bank comments and changes  ·  19 August 2026",
        size=8,
        color=MUTED,
    )


def cover(doc) -> None:
    bar = doc.add_table(rows=1, cols=1)
    c = bar.cell(0, 0)
    shade(c, COVER)
    no_borders(c)
    cell_margins(c, top=140, bottom=140, left=160, right=160)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(), "KP-RAP", size=11, bold=True, color=WHITE)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(0)
    set_run(p2.add_run(), "Khyber Pakhtunkhwa Rural Accessibility Project", size=9, color="E8EEF4")
    set_row_height(bar.rows[0], 900)

    gold = doc.add_table(rows=1, cols=1)
    g = gold.cell(0, 0)
    shade(g, GOLD)
    no_borders(g)
    g.paragraphs[0].paragraph_format.space_before = Pt(0)
    g.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_row_height(gold.rows[0], 120)

    P(doc, "", space_after=10)
    P(doc, "KHYBER PAKHTUNKHWA RURAL ACCESSIBILITY PROJECT", size=9, bold=True, color=GOLD, space_after=4)
    P(doc, "WORLD BANK COMMENTS AND CHANGES", size=20, bold=True, color=TEAL, space_before=2, space_after=4)
    P(
        doc,
        "Quality Assurance Workplan  ·  Version 2.0",
        size=12,
        color=INK,
        space_after=14,
    )

    meta_rows = [
        ("Related document", "RAP Quality Assurance Workplan 2.0 (19 August 2026)"),
        ("Prepared by", "Alliance of Excellence (AoE)"),
        ("Date", "19 August 2026"),
        ("What this is", "Each World Bank comment, and the change made in the workplan"),
    ]
    meta = doc.add_table(rows=len(meta_rows), cols=2)
    meta.autofit = False
    for i, (k, v) in enumerate(meta_rows):
        meta.cell(i, 0).width = Cm(5.2)
        meta.cell(i, 1).width = Cm(11.8)
        meta_row(meta, i, k, v)
    P(doc, "", space_after=8)


def build() -> list[str]:
    apply_palette("rap")
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    setup_header(doc)
    cover(doc)
    doc.add_page_break()

    h1(doc, "", "World Bank comments and our changes")
    body(
        doc,
        "The comments below were given on Version 1.0 of the quality-assurance workplan. "
        "The right-hand column is what was put in Version 2.0. AoE will follow Version 2.0 as written.",
    )
    tbl(
        doc,
        ["World Bank comment", "Change in Version 2.0"],
        ROWS,
        col_widths=[5.5, 6.5],
    )

    P(doc, "", space_after=12)
    P(
        doc,
        "Alliance of Excellence (AoE)  ·  KP-RAP  ·  World Bank comments and changes",
        size=8,
        color=MUTED,
        align="center",
    )

    dests = [
        ROOT / OUT_NAME,
        Path.home() / "OneDrive" / "Desktop" / OUT_NAME,
        Path.home() / "Desktop" / OUT_NAME,
    ]
    buf = BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    written = []
    for d in dests:
        try:
            d.parent.mkdir(parents=True, exist_ok=True)
            d.write_bytes(data)
            written.append(str(d))
        except Exception:
            pass
    return written


if __name__ == "__main__":
    paths = build()
    print("Wrote:")
    for p in paths:
        print(" ", p)
