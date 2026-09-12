"""
Shared builder for the "Ghalti ki Wajah aur Hal" Word documents.
Simple mixed Urdu-English field coaching notes, one per district,
companion to each district's Error Quality Report.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL_DEEP = RGBColor(0x0A, 0x4F, 0x47)
TEAL = RGBColor(0x0F, 0x6F, 0x63)
AMBER = RGBColor(0xA1, 0x5C, 0x12)
GREEN = RGBColor(0x2F, 0x7A, 0x44)
INK = RGBColor(0x1D, 0x26, 0x24)
INK_SOFT = RGBColor(0x4D, 0x5A, 0x56)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

AMBER_FILL = "FBF0E2"
GREEN_FILL = "ECF5EC"
TEAL_FILL = "E5F2EF"
HEADER_FILL = "0F6F63"
CHIP_FILL = "EEF1EF"


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def tbl_borders_light(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "DDE3E0")
        borders.append(el)
    for edge in ("insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_run(p, text, size=10.5, bold=False, color=INK, font="Calibri", italic=False):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r


def doc_heading(doc, district, summary_line, stats_line):
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = True
    cell = band.rows[0].cells[0]
    shade_cell(cell, HEADER_FILL)
    set_cell_margins(cell, top=260, bottom=260, left=280, right=280)
    no_borders(band)

    p1 = cell.paragraphs[0]
    add_run(p1, "KP-RAP PROJECT | FIELD COACHING NOTE", size=9, bold=True, color=WHITE)

    p2 = cell.add_paragraph()
    p2.space_before = Pt(4)
    add_run(p2, "Ghalti ki Wajah aur Hal", size=22, bold=True, color=WHITE, font="Cambria")

    p3 = cell.add_paragraph()
    p3.space_before = Pt(4)
    add_run(p3, summary_line, size=10.5, color=WHITE)

    p4 = cell.add_paragraph()
    p4.space_before = Pt(6)
    add_run(p4, stats_line, size=9, bold=True, color=WHITE)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def intro_box(doc):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shade_cell(cell, "F6F7F5")
    set_cell_margins(cell, top=200, bottom=200, left=220, right=220)
    tbl_borders_light(tbl)
    p = cell.paragraphs[0]
    add_run(p, "Har rule ke 2 parts hain: ", size=10.5, color=INK_SOFT)
    add_run(p, "Wajah", size=10.5, bold=True, color=INK)
    add_run(p, " (yeh error kyun hota hai) aur ", size=10.5, color=INK_SOFT)
    add_run(p, "Hal", size=10.5, bold=True, color=INK)
    add_run(
        p,
        " (ise kaise rokna hai). Supervisor is note ko morning briefing mein enumerators ke saath "
        "parh kar discuss karein.",
        size=10.5,
        color=INK_SOFT,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def section_heading(doc, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    add_run(p, title, size=15, bold=True, color=TEAL_DEEP, font="Cambria")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    add_run(p2, subtitle, size=9.5, italic=True, color=INK_SOFT)
    pBorder = doc.add_paragraph()
    pBorder.paragraph_format.space_after = Pt(6)
    pPr = pBorder._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "DDE3E0")
    bottom.set(qn("w:space"), "1")
    pbdr.append(bottom)
    pPr.append(pbdr)


def rule_card(doc, code, title, count, wajah, hal):
    head = doc.add_table(rows=1, cols=3)
    head.autofit = True
    head.columns[0].width = Cm(7.2)
    head.columns[1].width = Cm(6.2)
    head.columns[2].width = Cm(2.0)
    no_borders(head)

    c0, c1, c2 = head.rows[0].cells
    set_cell_margins(c0, top=40, bottom=40, left=80, right=80)
    set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
    set_cell_margins(c2, top=40, bottom=40, left=80, right=80)

    shade_cell(c0, TEAL_FILL)
    p0 = c0.paragraphs[0]
    add_run(p0, code, size=8, bold=True, color=TEAL_DEEP, font="Consolas")

    p1 = c1.paragraphs[0]
    add_run(p1, title, size=10.5, bold=True, color=INK)

    shade_cell(c2, CHIP_FILL)
    c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(c2.paragraphs[0], count, size=8.5, bold=True, color=INK_SOFT, font="Consolas")

    body = doc.add_table(rows=1, cols=2)
    body.autofit = True
    body.columns[0].width = Cm(7.7)
    body.columns[1].width = Cm(7.7)
    tbl_borders_light(body)

    wc, hc = body.rows[0].cells
    shade_cell(wc, AMBER_FILL)
    shade_cell(hc, GREEN_FILL)
    set_cell_margins(wc, top=140, bottom=140, left=180, right=180)
    set_cell_margins(hc, top=140, bottom=140, left=180, right=180)

    wp_label = wc.paragraphs[0]
    add_run(wp_label, "WAJAH", size=8.5, bold=True, color=AMBER)
    wp = wc.add_paragraph()
    add_run(wp, wajah, size=9.5, color=INK)

    hp_label = hc.paragraphs[0]
    add_run(hp_label, "HAL", size=8.5, bold=True, color=GREEN)
    hp = hc.add_paragraph()
    add_run(hp, hal, size=9.5, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def enumerator_section(doc, enumerators):
    n = len(enumerators)
    section_heading(
        doc,
        "Enumerator Coaching",
        f"In {n} enumerators ka score sab se kam hai. In ko yeh guidance khas tor par dohrayein",
    )

    tbl = doc.add_table(rows=1, cols=3)
    tbl.autofit = True
    tbl.columns[0].width = Cm(3.4)
    tbl.columns[1].width = Cm(5.5)
    tbl.columns[2].width = Cm(6.5)
    tbl_borders_light(tbl)

    hdr = tbl.rows[0].cells
    for i, txt in enumerate(["Enumerator", "Sabse Zyada Masla", "Kaise Bachein"]):
        shade_cell(hdr[i], HEADER_FILL)
        set_cell_margins(hdr[i], top=80, bottom=80, left=140, right=140)
        add_run(hdr[i].paragraphs[0], txt, size=9.5, bold=True, color=WHITE)

    for name, score, mistake, fix in enumerators:
        row = tbl.add_row().cells
        for c in row:
            set_cell_margins(c, top=100, bottom=100, left=140, right=140)
        p0 = row[0].paragraphs[0]
        add_run(p0, name, size=9.5, bold=True, color=INK)
        p0b = row[0].add_paragraph()
        add_run(p0b, score, size=8, color=INK_SOFT, font="Consolas")

        add_run(row[1].paragraphs[0], mistake, size=9.5, color=INK)
        add_run(row[2].paragraphs[0], fix, size=9.5, color=INK)


def footer_note(doc, district):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:color"), "DDE3E0")
    top.set(qn("w:space"), "4")
    pbdr.append(top)
    pPr.append(pbdr)
    add_run(p, f"KP-RAP M&E  |  Field Coaching Note  |  {district}", size=8.5, color=INK_SOFT)
    p2 = doc.add_paragraph()
    add_run(p2, "Source: Error Quality Report, 10-Sep-2026", size=8.5, color=INK_SOFT)


def build_document(
    district,
    summary_line,
    stats_line,
    critical_rules,
    quality_rules,
    enumerators,
    out_path,
):
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    doc_heading(doc, district, summary_line, stats_line)
    intro_box(doc)

    section_heading(
        doc,
        "Critical Issues",
        "Yeh ghaltiyan tracking aur household integrity ko seedha nuqsan pohanchati hain",
    )
    for code, title, count, wajah, hal in critical_rules:
        rule_card(doc, code, title, count, wajah, hal)

    section_heading(
        doc,
        "Quality Flags",
        "Yeh chhoti ghaltiyan hain lekin data ki quality kamzor karti hain",
    )
    for code, title, count, wajah, hal in quality_rules:
        rule_card(doc, code, title, count, wajah, hal)

    enumerator_section(doc, enumerators)
    footer_note(doc, district)

    doc.save(out_path)
    print(f"Saved: {out_path}")
